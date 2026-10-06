from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT = Path(__file__).resolve().parents[1] / "MCP_Tool_Test_Matrix.xlsx"
NAVY = "172033"
BLUE = "087EA4"
PALE_BLUE = "E8F4F8"
PALE_GREEN = "E2F3E9"
PALE_RED = "FCE8E6"
WHITE = "FFFFFF"

TOOLS = [
    ("site_metadata_summary", "Site name, ID, buildings, device and meter counts", "What site am I viewing? Show its name, ID, building count, device count, and meter count.", "Site identity and counts only."),
    ("site_device_list", "Device inventory and status", "List this site's devices with IDs, online status, type, and last-seen time when available.", "Device rows plus status totals or a clear result limit."),
    ("active_alarm_summary", "Currently active alarms", "How many active alarms are on this site now? Group them by severity and show the latest five.", "Active totals, severity, and latest alarm details."),
    ("meter_status_summary", "Online, offline, stale, unknown meter status", "Which meters are offline, stale, or not reporting at this site? Include last-seen times where available.", "Status counts and affected meters; missing fields are stated plainly."),
    ("site_energy_summary", "Site energy across a time range", "Show daily site energy for the last 7 days, with dates, values, and units.", "Daily values for the requested window; chart only when supported and enabled."),
    ("telemetry_top_consumers", "Rank devices by energy use", "Which five devices consumed the most energy at this site in the last 7 days? Include values and units.", "Ranked devices with comparable values and an explicit time window."),
    ("alarm_frequency_summary", "Historical alarm frequency", "What alarm types occurred most often at this site in the last 7 days? Rank the top five.", "Alarm categories/counts over the requested window."),
    ("telemetry_timeseries", "Time-series for a device/metric", "Show the voltage time series for device ID 380 for the last 24 hours.", "Device and metric are correct; timestamps, units, and time window are clear."),
    ("energy_comparison_summary", "Compare two energy periods", "Compare this week with last week for site energy. Show both totals and the difference.", "Both periods, values, units, and direction of change."),
    ("data_availability_summary", "Coverage and missing readings", "Was energy data complete for this site over the last 7 days? Identify missing periods if available.", "Coverage/missing intervals; do not invent completeness."),
    ("alarm_detail_lookup", "Look up a specific alarm", "Show details for alarm ID 123, including device, severity, start time, and current status.", "Replace 123 with a real alarm ID before judging data validity."),
    ("power_quality_summary", "Power-quality events/metrics", "Summarize power-quality issues at this site for the last 7 days, including sag, swell, THD, or power factor when available.", "Only returned metrics/events, with units and coverage limits."),
    ("demand_peak_summary", "Maximum/peak demand", "What was the highest demand at this site today? Give the value, unit, timestamp, and source meter if returned.", "Demand peak, period, timestamp and source where available; clarify missing scope."),
    ("tariff_cost_summary", "Energy cost under configured tariffs", "Show this site's electricity cost this month, with the tariff and energy values used if available.", "Cost and currency only when returned; identify absent tariff inputs."),
    ("device_energy_breakdown", "Energy contribution by device", "Break down this site's energy by device for the last 7 days and rank the largest contributors.", "Device-level values over one consistent period; totals reconcile when possible."),
    ("energy_forecast", "Forecast from historical energy data", "Forecast daily energy use for this site for the next 7 days and state the method or limitations.", "Clearly labeled forecast, horizon and uncertainty/limitations."),
    ("anomaly_detection_summary", "Unusual energy behavior", "Find unusual energy usage at this site over the last 7 days and show the affected dates/devices and baseline if available.", "Evidence and baseline are shown; no unsupported cause claims."),
    ("report_summary", "Combine operational summaries into a report", "Prepare a weekly EMS report for this site covering energy, top consumers, alarms, offline meters, and demand.", "Each requested section is present or explicitly marked unavailable."),
]

MULTI = [
    ("Site health", "Give me a site health summary for the last 7 days: active alarms, offline meters, max demand, top energy consumers, and unusual energy use.", "active_alarm_summary; meter_status_summary; demand_peak_summary; telemetry_top_consumers; anomaly_detection_summary", "Every requested section has its own source result or an explicit unavailable note."),
    ("Demand investigation", "Show today's maximum site demand with time and source meter. Then list the candidate devices to inspect, keeping measured demand separate from device inventory.", "demand_peak_summary; site_device_list", "Does not imply listed devices caused the peak unless data links them."),
    ("Energy and alarms", "Compare site energy this week with last week, then summarize the most frequent alarms over the same periods.", "energy_comparison_summary; alarm_frequency_summary", "Both requested analyses run and use aligned periods."),
    ("Data-quality check", "For the last 7 days, compare daily energy totals with data coverage and list dates where readings are missing or suspicious.", "site_energy_summary; data_availability_summary; anomaly_detection_summary", "Each claim ties to returned data; gaps remain distinct from anomalies."),
    ("Operator report", "Create an operator report for the last 7 days with energy totals, top 5 consumers, critical alarms, offline meters, and practical checks supported by the data.", "site_energy_summary; telemetry_top_consumers; active_alarm_summary; meter_status_summary", "No unsupported diagnosis; each action is traceable to a finding."),
]

MODELS = [
    ("Energy summary", "Show daily site energy for the last 7 days and explain the largest day-to-day change.", "site_energy_summary", "Uses the same MCP result and time range for every model."),
    ("Top consumers", "Rank the top five energy-consuming devices for this site over the last 7 days. State units and data gaps.", "telemetry_top_consumers", "No invented devices or values; ordering matches data."),
    ("Max demand", "What is the max demand today? Include value, unit, timestamp, and source meter only when returned.", "demand_peak_summary", "Correctly handles missing demand fields and avoids substituting device counts."),
    ("Multiple asks", "Compare energy this week versus last week, then tell me whether missing data could affect the comparison.", "energy_comparison_summary; data_availability_summary", "Answers both parts and separates facts from limitations."),
    ("Vague device follow-up", "Can you get max demand for the device?", "demand_peak_summary; then clarification/device selection if device is unspecified", "Asks for a device choice instead of guessing."),
]

DIRECT_PROOF = [
    ("List available MCP tools", "tools/list", "{}", "Confirms whether the deployed DaxView MCP exposes all expected tools, including voltage/current-capable telemetry."),
    ("Direct site metadata", "site_metadata_summary", '{"authorization_id":"<auth>","site_id":<site_id>}', "Proves read access and site scope before testing historical metrics."),
    ("Direct demand peak", "demand_peak_summary", '{"authorization_id":"<auth>","site_id":<site_id>,"start_time":"<iso>","end_time":"<iso>","timezone":"Asia/Kuala_Lumpur"}', "If this returns NO_DATA, test the telemetry fallback below."),
    ("Manual demand fallback source", "telemetry_timeseries", '{"authorization_id":"<auth>","site_id":<site_id>,"metric":"demand","start_time":"<iso>","end_time":"<iso>","bucket":"1h","aggregation":"auto","limit":1000}', "Use max(value/demand_kw/kw) from returned rows as fallback evidence."),
    ("Voltage availability", "telemetry_timeseries", '{"authorization_id":"<auth>","device_id":<device_id>,"metric":"voltage","start_time":"<iso>","end_time":"<iso>","bucket":"1h","aggregation":"auto","limit":500}', "Records whether voltage is accepted and whether rows contain voltage values."),
    ("Current availability", "telemetry_timeseries", '{"authorization_id":"<auth>","device_id":<device_id>,"metric":"current","start_time":"<iso>","end_time":"<iso>","bucket":"1h","aggregation":"auto","limit":500}', "Records whether current is accepted and whether rows contain current values."),
    ("Metric coverage check", "data_availability_summary", '{"authorization_id":"<auth>","site_id":<site_id>,"metric":"voltage","start_time":"<iso>","end_time":"<iso>","timezone":"Asia/Kuala_Lumpur"}', "If supported, proves missing voltage data separately from unsupported telemetry."),
]


def add_table_sheet(wb, title, headers, rows, widths, table_name):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + index)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row[0].row].height = 54
    if ws.max_row > 1:
        table = Table(displayName=table_name, ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)
    return ws


def add_list_validation(ws, column, values, start=2, end=300):
    validation = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    ws.add_data_validation(validation)
    validation.add(f"{column}{start}:{column}{end}")


def main():
    wb = Workbook()
    intro = wb.active
    intro.title = "Start Here"
    rows = [
        ("Purpose", "Use this workbook to verify routing, MCP data, answer grounding, model quality, and chart output."),
        ("Important", "The operation names in the test matrix are intended expectations. Confirm each tool exists in the deployed MCP tools/list before testing; an AI-server allowlist does not prove the remote MCP implements it."),
        ("Run order", "1) Select a site and time range in DaxView. 2) Ask one independent question. 3) Open the AI MCP Trace Dashboard. 4) Record the selected tool, normalized arguments, MCP result, model answer, and chart preview. 5) Mark pass/fail in the workbook."),
        ("Expected trace", "daxview_tool_selection_debug -> daxview_data_plan_request -> mcp_tool_request -> mcp_tool_response -> mcp_answer_refine_request/response -> ai_final_response_debug."),
        ("Payload visibility", "MCP arguments and returned payloads appear in the debug dashboard only when DAXVIEW_MCP_DEBUG_RESPONSE=true. The dashboard must also be enabled and protected with AI_DEBUG_DASHBOARD_KEY."),
        ("Model comparison", "The configured CHAT_MODEL remains primary. Compare Qwen3:8b or DeepSeek against the same MCP result. Score each independently; don't judge from wording alone."),
        ("Scoring", "Grounding: values/devices match data. Completeness: all parts answered. Tool routing: expected tool(s) called. Clarity: readable and appropriately qualified. Use 1 (poor) to 5 (excellent)."),
        ("Chart check", "A chart is expected only for compatible returned data and when AI_CHARTS_ENABLED=true. No chart for empty/non-numeric results is a correct outcome."),
        ("Privacy", "Debug snapshots may contain operational site data. Keep dashboard access restricted; do not paste keys, assertions, or authorization IDs into this workbook."),
    ]
    intro.append(["Topic", "Instructions"])
    for row in rows:
        intro.append(row)
    intro.freeze_panes = "A2"
    intro.column_dimensions["A"].width = 24
    intro.column_dimensions["B"].width = 118
    for cell in intro[1]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
    for row in intro.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        intro.row_dimensions[row[0].row].height = 44
        row[0].font = Font(bold=True, color=BLUE)

    independent_rows = [
        (index, tool, purpose, prompt, tool, expected, "Confirm using tools/list", "Not tested", "", "", "", "")
        for index, (tool, purpose, prompt, expected) in enumerate(TOOLS, start=1)
    ]
    ws = add_table_sheet(
        wb, "Tool Tests",
        ["#", "Expected tool", "What it should answer", "Test question", "Expected operation", "Expected answer checks", "MCP availability", "Status", "Actual operation(s)", "Data returned?", "Chart?", "Notes / issue"],
        independent_rows, [6, 30, 34, 76, 30, 56, 28, 18, 34, 18, 14, 42], "ToolTests",
    )
    add_list_validation(ws, "H", ["Not tested", "Pass", "Fail", "Blocked", "Tool unavailable"])
    add_list_validation(ws, "J", ["Yes", "No", "Partial"])
    add_list_validation(ws, "K", ["Yes", "No", "N/A"])

    multi_rows = [(name, prompt, expected_tools, acceptance, "Not tested", "", "", "") for name, prompt, expected_tools, acceptance in MULTI]
    ws = add_table_sheet(
        wb, "Multi Tool Tests",
        ["Scenario", "Test question", "Expected operation(s)", "Pass criteria", "Status", "Actual operation(s)", "All parts answered?", "Notes"],
        multi_rows, [24, 86, 66, 62, 18, 54, 22, 42], "MultiToolTests",
    )
    add_list_validation(ws, "E", ["Not tested", "Pass", "Fail", "Blocked"])
    add_list_validation(ws, "G", ["Yes", "No", "Partial"])

    model_rows = [(name, prompt, expected, checks, "CHAT_MODEL (current)", "", "", "", "qwen3:8b", "", "", "") for name, prompt, expected, checks in MODELS]
    ws = add_table_sheet(
        wb, "Model Comparison",
        ["Scenario", "Shared prompt", "Expected tool(s)", "What to judge", "Primary model", "Primary grounding (1-5)", "Primary completeness (1-5)", "Primary clarity (1-5)", "Stronger model", "Stronger grounding (1-5)", "Stronger completeness (1-5)", "Stronger clarity (1-5)"],
        model_rows, [24, 82, 48, 66, 22, 24, 26, 22, 22, 26, 28, 22], "ModelComparison",
    )
    for column in ("F", "G", "H", "J", "K", "L"):
        add_list_validation(ws, column, ["1", "2", "3", "4", "5"])

    proof_rows = [(name, method, payload, evidence, "Not tested", "", "", "", "") for name, method, payload, evidence in DIRECT_PROOF]
    ws = add_table_sheet(
        wb, "Direct MCP Proof",
        ["Proof case", "MCP method/tool", "Payload template", "Evidence to capture", "Status", "Accepted?", "Rows/data?", "Error/code", "Notes / trace ID"],
        proof_rows, [28, 28, 82, 58, 18, 16, 16, 28, 48], "DirectMcpProof",
    )
    add_list_validation(ws, "E", ["Not tested", "Pass", "Fail", "Blocked", "Tool unavailable"])
    add_list_validation(ws, "F", ["Yes", "No", "N/A"])
    add_list_validation(ws, "G", ["Yes", "No", "Partial", "N/A"])

    log_headers = ["Date/time", "Tester", "Question", "Expected tool(s)", "Actual tool(s)", "Arguments correct?", "MCP status", "Data result", "Primary model", "Primary score (1-5)", "Compare model", "Compare score (1-5)", "Chart expected?", "Chart shown?", "Latency (s)", "Overall result", "Notes / trace ID"]
    ws = add_table_sheet(wb, "Run Log", log_headers, [[""] * len(log_headers) for _ in range(30)], [22, 20, 68, 42, 42, 20, 18, 18, 22, 20, 22, 20, 18, 18, 16, 18, 48], "RunLog")
    for column, values in (("F", ["Yes", "No", "Partial"]), ("G", ["Success", "Error", "Denied", "Timeout"]), ("H", ["Yes", "No", "Partial"]), ("M", ["Yes", "No", "N/A"]), ("N", ["Yes", "No", "N/A"]), ("P", ["Pass", "Fail", "Blocked"])):
        add_list_validation(ws, column, values, end=301)
    for column in ("J", "L"):
        add_list_validation(ws, column, ["1", "2", "3", "4", "5"], end=301)
    ws.conditional_formatting.add("P2:P301", CellIsRule(operator="equal", formula=['"Pass"'], fill=PatternFill("solid", fgColor=PALE_GREEN)))
    ws.conditional_formatting.add("P2:P301", CellIsRule(operator="equal", formula=['"Fail"'], fill=PatternFill("solid", fgColor=PALE_RED)))

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.tabColor = BLUE if sheet.title != "Run Log" else "16A34A"
    wb.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()
